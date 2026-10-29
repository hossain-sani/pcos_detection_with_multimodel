"""Step 3: Tri-Stage Hybrid Feature Selection (63 -> 30 -> 18 -> 16).

Proposed methodology, fit ONLY on Step-2 balanced holdout train (zero leakage):
  Stage 1: Chi2 + ANOVA rank-average filter      63 -> 30
  Stage 2: Boruta (shadow features) + CatBoost RFE rank-merge  30 -> 18
  Stage 3: ElasticNet stability selection (50 repeats)          18 -> 16
  Wrapper RFE (LR/XGB/RF/CB) on the 30 Stage-1 features for Table 3.

Single-run script: executing this file regenerates all Step-3 outputs
(updated data CSVs for every stage, figures, QC reports, filled T3+T4 docx).
"""

from __future__ import annotations

import json
import shutil
import sys
import warnings
from datetime import datetime, timezone
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from boruta.boruta_py import BorutaPy
from catboost import CatBoostClassifier
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from lightgbm import LGBMClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn.feature_selection import RFE, chi2, f_classif
from sklearn.linear_model import LogisticRegression
from sklearn.preprocessing import StandardScaler
from xgboost import XGBClassifier

warnings.filterwarnings("ignore")

SEED = 42
N_STAGE1 = 30
N_STAGE2 = 18
N_STAGE3 = 16
STABILITY_REPEATS = 50
STABILITY_SUBSAMPLE = 0.8
STABILITY_THRESHOLD = 0.7
RFE_TOP_KS = [20, 16, 12]

ROOT = Path(__file__).resolve().parents[2]
STEP1_DIR = ROOT / "step1_preprocessing"
STEP2_DIR = ROOT / "step2_balancing"
STEP_DIR = ROOT / "step3_feature_selection"
RESULTS_DIR = STEP_DIR / "results"
DATA_DIR = RESULTS_DIR / "data"
ARTIFACTS_DIR = RESULTS_DIR / "artifacts"
QC_DIR = RESULTS_DIR / "qc_reports"
FIG_DIR = RESULTS_DIR / "figures"
TABLES_DIR = STEP_DIR / "tables"

STEP1_RAW_CSV = STEP1_DIR / "results" / "pcos_cleaned_raw.csv"
STEP1_MANIFEST = STEP1_DIR / "results" / "artifacts" / "feature_manifest.json"
STEP2_BAL_X = STEP2_DIR / "results" / "balanced_folds" / "holdout_train_balanced.csv"
STEP2_BAL_Y = STEP2_DIR / "results" / "balanced_folds" / "holdout_train_balanced_target.csv"
STEP2_DOCX = STEP2_DIR / "tables" / "Step2_Tables_Filled.docx"
TEMPLATE_DOCX = ROOT / "docs" / "All_Tables_Templates.docx"
OUT_DOCX = TABLES_DIR / "Step3_Tables_Filled.docx"

TARGET = "PCOS (Y/N)"
META_COLS = ["Patient_File_No", "Split"]

DOCX_FONT = "Calibri"
DOCX_SIZE = 8


def log(msg: str) -> None:
    print(f"[step3] {msg}", flush=True)


def ensure_dirs() -> None:
    for path in (DATA_DIR, ARTIFACTS_DIR, QC_DIR, FIG_DIR, TABLES_DIR):
        path.mkdir(parents=True, exist_ok=True)


def load_inputs() -> tuple[pd.DataFrame, pd.DataFrame, pd.Series, list[str], dict]:
    full = pd.read_csv(STEP1_RAW_CSV)
    manifest = json.loads(STEP1_MANIFEST.read_text(encoding="utf-8"))
    feature_cols = manifest["feature_columns"]
    bal_X = pd.read_csv(STEP2_BAL_X)
    bal_y = pd.read_csv(STEP2_BAL_Y)[TARGET].astype(int)
    missing = [c for c in feature_cols if c not in full.columns or c not in bal_X.columns]
    if missing:
        raise SystemExit(f"Missing feature columns: {missing}")
    log(f"Full data: {full.shape[0]} x {len(feature_cols)} features | balanced train: {bal_X.shape[0]}")
    return full, bal_X, bal_y, feature_cols, manifest


def stage1_filter(bal_X: pd.DataFrame, bal_y: pd.Series, feature_cols: list[str]) -> pd.DataFrame:
    X_nn = bal_X[feature_cols].clip(lower=0)
    chi2_stat, chi2_p = chi2(X_nn, bal_y)
    f_stat, f_p = f_classif(X_nn.replace([np.inf, -np.inf], np.nan).fillna(X_nn.median()), bal_y)
    report = pd.DataFrame({
        "feature": feature_cols,
        "chi2_stat": chi2_stat, "chi2_p": chi2_p,
        "anova_F": f_stat, "anova_p": f_p,
    })
    report["chi2_rank"] = report["chi2_stat"].rank(ascending=False, method="average")
    report["anova_rank"] = report["anova_F"].rank(ascending=False, method="average")
    report["combined_rank"] = (report["chi2_rank"] + report["anova_rank"]) / 2.0
    report = report.sort_values(
        ["combined_rank", "chi2_stat", "feature"],
        ascending=[True, False, True],
        na_position="last",
    ).reset_index(drop=True)
    report["stage1_rank"] = np.arange(1, len(report) + 1)
    report["stage1_keep"] = report["stage1_rank"] <= N_STAGE1
    log(f"Stage 1: Chi2+ANOVA -> {N_STAGE1} features")
    return report


def stage2_select(bal_X: pd.DataFrame, bal_y: pd.Series, stage1_feats: list[str]) -> tuple[pd.DataFrame, list[str]]:
    X = bal_X[stage1_feats].to_numpy(dtype=float)
    lgbm = LGBMClassifier(n_estimators=300, random_state=SEED, n_jobs=-1, verbose=-1)
    boruta = BorutaPy(lgbm, n_estimators=300, random_state=SEED, max_iter=100)
    boruta.fit(X, bal_y.to_numpy())
    boruta_df = pd.DataFrame({
        "feature": stage1_feats,
        "boruta_support": boruta.support_.tolist(),
        "boruta_weak": boruta.support_weak_.tolist(),
        "boruta_rank": boruta.ranking_.tolist(),
    })
    log(f"Stage 2a: Boruta confirmed {int(boruta.support_.sum())}/{len(stage1_feats)}")

    cb = CatBoostClassifier(iterations=300, depth=4, learning_rate=0.1,
                            random_seed=SEED, verbose=0, allow_writing_files=False)
    rfe = RFE(cb, n_features_to_select=1, step=1, importance_getter="feature_importances_")
    rfe.fit(X, bal_y.to_numpy())
    rfe_df = pd.DataFrame({
        "feature": stage1_feats,
        "cb_rfe_rank": rfe.ranking_.tolist(),
        "cb_rfe_selected": rfe.support_.tolist(),
    })
    log("Stage 2b: CatBoost RFE ranking done")

    merged = boruta_df.merge(rfe_df, on="feature")
    merged["boruta_rank_norm"] = merged["boruta_rank"].rank(pct=True)
    merged["rfe_rank_norm"] = merged["cb_rfe_rank"].rank(ascending=False, pct=True)
    merged["stage2_score"] = (merged["boruta_rank_norm"] + merged["rfe_rank_norm"]) / 2.0
    merged = merged.sort_values(["stage2_score", "cb_rfe_rank"]).reset_index(drop=True)
    merged["stage2_rank"] = np.arange(1, len(merged) + 1)
    merged["stage2_keep"] = merged["stage2_rank"] <= N_STAGE2
    top = merged.loc[merged["stage2_keep"], "feature"].tolist()
    log(f"Stage 2: Boruta + CatBoost RFE -> {len(top)} features")
    return merged, top


def stage3_stability(bal_X: pd.DataFrame, bal_y: pd.Series, stage2_feats: list[str]) -> pd.DataFrame:
    rng = np.random.default_rng(SEED)
    X = bal_X[stage2_feats].to_numpy(dtype=float)
    y = bal_y.to_numpy()
    n = len(y)
    counts = dict.fromkeys(stage2_feats, 0)
    coef_abs = dict.fromkeys(stage2_feats, 0.0)
    scaler = StandardScaler()
    Xs = scaler.fit_transform(X)
    for rep in range(STABILITY_REPEATS):
        sub = rng.choice(n, size=int(n * STABILITY_SUBSAMPLE), replace=False)
        y_sub = y[sub]
        if len(np.unique(y_sub)) < 2:
            continue
        lr = LogisticRegression(
            penalty="elasticnet", solver="saga", l1_ratio=0.5, C=1.0,
            max_iter=5000, random_state=SEED + rep,
        )
        lr.fit(Xs[sub], y_sub)
        for j, feat in enumerate(stage2_feats):
            if lr.coef_[0, j] != 0:
                counts[feat] += 1
            coef_abs[feat] += abs(float(lr.coef_[0, j]))
    df = pd.DataFrame({
        "feature": stage2_feats,
        "selection_frequency": [counts[f] / STABILITY_REPEATS for f in stage2_feats],
        "mean_abs_coef": [coef_abs[f] / STABILITY_REPEATS for f in stage2_feats],
    })
    df = df.sort_values(["selection_frequency", "mean_abs_coef"], ascending=False).reset_index(drop=True)
    df["stability_rank"] = np.arange(1, len(df) + 1)
    df["stage3_keep"] = df["stability_rank"] <= N_STAGE3
    n_above = int((df["selection_frequency"] >= STABILITY_THRESHOLD).sum())
    log(f"Stage 3: ElasticNet stability ({STABILITY_REPEATS}x) -> {N_STAGE3} features "
        f"({n_above} above {STABILITY_THRESHOLD:.0%} threshold)")
    return df


def wrapper_rfe_table3(bal_X: pd.DataFrame, bal_y: pd.Series, stage1_feats: list[str]) -> pd.DataFrame:
    X = bal_X[stage1_feats].to_numpy(dtype=float)
    estimators = {
        "CB": CatBoostClassifier(iterations=200, depth=4, random_seed=SEED,
                                 verbose=0, allow_writing_files=False),
        "XGB": XGBClassifier(n_estimators=200, max_depth=4, learning_rate=0.1,
                             random_state=SEED, n_jobs=-1, eval_metric="logloss",
                             verbosity=0),
        "LR": LogisticRegression(max_iter=5000, random_state=SEED),
        "RF": RandomForestClassifier(n_estimators=200, random_state=SEED, n_jobs=-1),
    }
    rows = []
    for name, est in estimators.items():
        getter = "auto" if name == "LR" else "feature_importances_"
        rfe = RFE(est, n_features_to_select=1, step=1, importance_getter=getter)
        rfe.fit(X, bal_y.to_numpy())
        rows.append(pd.DataFrame({
            "feature": stage1_feats,
            "model": name,
            "rank": rfe.ranking_.tolist(),
        }))
    out = pd.concat(rows, ignore_index=True)
    for k in RFE_TOP_KS:
        out[f"top{k}"] = out["rank"] <= k
    log(f"Wrapper RFE done for Table 3: models {sorted(estimators)} x top{RFE_TOP_KS}")
    return out


def save_updated_data(full: pd.DataFrame, manifest: dict,
                      stage1: list[str], stage2: list[str], stage3: list[str]) -> None:
    for tag, feats in [("stage1_chi2_anova_30", stage1),
                       ("stage2_boruta_cb_rfe_18", stage2),
                       ("stage3_final_16", stage3)]:
        cols = [c for c in META_COLS if c in full.columns] + [TARGET] + feats
        out = full[cols].copy()
        path = DATA_DIR / f"{tag}.csv"
        out.to_csv(path, index=False)
        log(f"Saved updated data: {path.name} ({out.shape[0]} rows x {out.shape[1]} cols)")


def make_figures(s1: pd.DataFrame, s2: pd.DataFrame, s3: pd.DataFrame,
                 wrapper: pd.DataFrame, n_total: int) -> None:
    sns.set_theme(style="whitegrid", context="notebook")

    fig, ax = plt.subplots(figsize=(7, 4.2))
    stages = ["Input", "Stage 1\nChi2+ANOVA", "Stage 2\nBoruta+CB-RFE", "Stage 3\nElasticNet"]
    counts = [n_total, N_STAGE1, N_STAGE2, N_STAGE3]
    sns.barplot(x=stages, y=counts, hue=stages,
                palette=["#8da0cb", "#66c2a5", "#fc8d62", "#e78ac3"],
                legend=False, ax=ax)
    for i, v in enumerate(counts):
        ax.text(i, v + 1.2, str(v), ha="center", fontsize=11, fontweight="bold")
    ax.set_ylabel("Number of features")
    ax.set_title("Tri-stage hybrid feature selection funnel")
    ax.set_ylim(0, n_total * 1.12)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "selection_funnel.png", dpi=150)
    plt.close(fig)

    top30 = s1[s1["stage1_keep"]].sort_values("combined_rank", ascending=True)
    final_set = set(s3.loc[s3["stage3_keep"], "feature"])
    stage2_set = set(s2.loc[s2["stage2_keep"], "feature"])
    fig, ax = plt.subplots(figsize=(9, 9))
    colors = ["#e78ac3" if f in final_set else "#a6d854"
              if f in stage2_set else "#8da0cb"
              for f in top30["feature"]]
    y_pos = np.arange(len(top30))
    ax.barh(y_pos, top30["combined_rank"].values, color=colors)
    ax.set_yticks(y_pos)
    ax.set_yticklabels([manifest_display(f) for f in top30["feature"]], fontsize=8)
    ax.invert_yaxis()
    ax.set_xlabel("Mean rank (Chi2, ANOVA) - lower is better")
    ax.set_title("Stage 1: top-30 features (pink=final 16, green=stage-2 only, blue=stage-1 only)")
    fig.tight_layout()
    fig.savefig(FIG_DIR / "stage1_top30_ranks.png", dpi=150)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(9, 5.5))
    s3_sorted = s3.sort_values("selection_frequency", ascending=True)
    bar_colors = ["#e78ac3" if k else "#8da0cb" for k in s3_sorted["stage3_keep"]]
    ax.barh(np.arange(len(s3_sorted)), s3_sorted["selection_frequency"], color=bar_colors)
    ax.axvline(STABILITY_THRESHOLD, color="#c44e52", ls="--", lw=1.2,
               label=f"threshold {STABILITY_THRESHOLD:.0%}")
    ax.set_yticks(np.arange(len(s3_sorted)))
    ax.set_yticklabels([manifest_display(f) for f in s3_sorted["feature"]], fontsize=8)
    ax.set_xlabel("Selection frequency (ElasticNet, 50 repeats)")
    ax.set_title("Stage 3: stability selection of 18 candidates")
    ax.legend()
    fig.tight_layout()
    fig.savefig(FIG_DIR / "stage3_stability_frequencies.png", dpi=150)
    plt.close(fig)

    pivot = wrapper.pivot(index="feature", columns="model", values="rank")
    pivot = pivot.reindex(s1.loc[s1["stage1_keep"], "feature"])
    fig, ax = plt.subplots(figsize=(7, 10))
    sns.heatmap(pivot, annot=True, fmt="g", cmap="YlGnBu_r", linewidths=0.4, ax=ax,
                cbar_kws={"label": "RFE rank (1=best)"})
    ax.set_title("Wrapper RFE ranks on Stage-1 features (Table 3 basis)")
    ax.set_xticklabels(ax.get_xticklabels(), rotation=0)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "wrapper_rfe_heatmap.png", dpi=150)
    plt.close(fig)


def manifest_display(feat: str) -> str:
    global _DISPLAY
    return _DISPLAY.get(feat, feat)


_DISPLAY: dict = {}


def _docx_set_cell(cell, text: str, align: str = "center") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    run.font.size = Pt(DOCX_SIZE)
    run.font.name = DOCX_FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), DOCX_FONT)


def build_tables_docx(stage1_feats: list[str], s1: pd.DataFrame, s2: pd.DataFrame,
                      s3: pd.DataFrame, wrapper: pd.DataFrame, final16: list[str]) -> None:
    src = STEP2_DOCX if STEP2_DOCX.exists() else TEMPLATE_DOCX
    shutil.copyfile(src, OUT_DOCX)
    doc = Document(OUT_DOCX)

    t3 = doc.tables[2]
    n_rows_data = len(t3.rows) - 3
    if len(stage1_feats) != n_rows_data:
        raise SystemExit(
            f"Table 3 row count mismatch: template has {n_rows_data} feature rows, "
            f"stage 1 selected {len(stage1_feats)} - cannot fill without editing format"
        )
    for i, feat in enumerate(stage1_feats):
        r = 3 + i
        _docx_set_cell(t3.cell(r, 0), manifest_display(feat), align="left")
        for k_i, k in enumerate(RFE_TOP_KS):
            for m_i, model in enumerate(["CB", "XGB", "LR", "RF"]):
                row = wrapper[(wrapper["feature"] == feat) & (wrapper["model"] == model)]
                mark = "✓" if bool(row[f"top{k}"].iloc[0]) else "–"
                _docx_set_cell(t3.cell(r, 1 + k_i * 4 + m_i), mark)
    log(f"Table 3 filled: {len(stage1_feats)} rows x 12 wrapper columns")

    t4 = doc.tables[3]
    if len(final16) != len(t4.rows) - 1:
        raise SystemExit("Table 4 row count mismatch")
    for i, feat in enumerate(final16):
        r = 1 + i
        _docx_set_cell(t4.cell(r, 1), manifest_display(feat), align="left")
    log(f"Table 4 filled: {len(final16)} final features")

    doc.save(OUT_DOCX)

    a, b = Document(TEMPLATE_DOCX), Document(OUT_DOCX)
    header_rows = {1: 1, 2: 3, 3: 3, 4: 1, 5: 1, 6: 1, 7: 1, 8: 3, 9: 3,
                   10: 2, 11: 2, 12: 1, 13: 1, 14: 2, 15: 1, 16: 1, 17: 1}
    ok = len(a.tables) == len(b.tables)
    for ti, (ta, tb) in enumerate(zip(a.tables, b.tables), start=1):
        for ri in range(header_rows.get(ti, 1)):
            for ci in range(len(ta.columns)):
                if ta.cell(ri, ci).text != tb.cell(ri, ci).text:
                    print(f"  header mismatch T{ti} R{ri}C{ci}")
                    ok = False
    log("Header verification: " + ("PASS - all headers byte-identical" if ok else "FAIL"))


def main() -> None:
    global _DISPLAY
    ensure_dirs()
    full, bal_X, bal_y, feature_cols, manifest = load_inputs()
    _DISPLAY = {k: v for k, v in manifest.get("display_names", {}).items()}
    _DISPLAY.update({b: b for b in manifest["biomarker_names"]})

    s1 = stage1_filter(bal_X, bal_y, feature_cols)
    s1.to_csv(QC_DIR / "stage1_chi2_anova_scores.csv", index=False, float_format="%.6g")
    stage1_feats = s1.loc[s1["stage1_keep"], "feature"].tolist()

    s2, stage2_feats = stage2_select(bal_X, bal_y, stage1_feats)
    s2.to_csv(QC_DIR / "stage2_boruta_rfe_merge.csv", index=False, float_format="%.6g")

    s3 = stage3_stability(bal_X, bal_y, stage2_feats)
    s3.to_csv(QC_DIR / "stability_frequencies.csv", index=False, float_format="%.6g")
    final16 = s3.loc[s3["stage3_keep"], "feature"].tolist()

    wrapper = wrapper_rfe_table3(bal_X, bal_y, stage1_feats)
    wrapper.to_csv(QC_DIR / "wrapper_rfe_ranks.csv", index=False)

    save_updated_data(full, manifest, stage1_feats, stage2_feats, final16)

    report = s1.merge(s2[["feature", "boruta_support", "boruta_rank", "cb_rfe_rank",
                          "stage2_rank", "stage2_keep"]], on="feature", how="left")
    report = report.merge(
        s3[["feature", "selection_frequency", "mean_abs_coef", "stability_rank", "stage3_keep"]],
        on="feature", how="left")
    report = report.sort_values("stage1_rank").reset_index(drop=True)
    report.to_csv(QC_DIR / "selection_report.csv", index=False, float_format="%.6g")

    make_figures(s1, s2, s3, wrapper, len(feature_cols))

    selected = {
        "stage1_chi2_anova_30": [manifest_display(f) for f in stage1_feats],
        "stage2_boruta_cb_rfe_18": [manifest_display(f) for f in stage2_feats],
        "stage3_final_16": [manifest_display(f) for f in final16],
        "internal_names": {"stage1": stage1_feats, "stage2": stage2_feats, "stage3": final16},
    }
    (ARTIFACTS_DIR / "selected_features.json").write_text(
        json.dumps(selected, indent=2), encoding="utf-8")

    build_tables_docx(stage1_feats, s1, s2, s3, wrapper, final16)

    summary = {
        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "protocol": {
            "fit_data": "Step-2 SMOTE-NC balanced holdout train (zero test leakage)",
            "stage1": f"Chi2 + ANOVA rank-average, top {N_STAGE1}",
            "stage2": f"Boruta(max_iter=100) + CatBoost-RFE rank-merge, top {N_STAGE2}",
            "stage3": f"ElasticNet stability, {STABILITY_REPEATS}x {STABILITY_SUBSAMPLE:.0%} subsample, "
                      f"l1_ratio=0.5, threshold {STABILITY_THRESHOLD}, top {N_STAGE3}",
            "wrapper_rfe": f"CB/XGB/LR/RF on stage-1 features, top {RFE_TOP_KS}",
        },
        "counts": {"input": len(feature_cols), "stage1": len(stage1_feats),
                   "stage2": len(stage2_feats), "stage3": len(final16)},
        "final_16": [manifest_display(f) for f in final16],
        "updated_data_files": sorted(p.name for p in DATA_DIR.glob("*.csv")),
    }
    (QC_DIR / "selection_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print()
    log("DONE")
    log(f"Funnel: {len(feature_cols)} -> {len(stage1_feats)} -> {len(stage2_feats)} -> {len(final16)}")
    log("Final 16: " + ", ".join(manifest_display(f) for f in final16))
    log(f"Updated data: {DATA_DIR.relative_to(ROOT)}")
    log(f"Docx: {OUT_DOCX.name} (T3+T4 filled)")


if __name__ == "__main__":
    sys.exit(main())
