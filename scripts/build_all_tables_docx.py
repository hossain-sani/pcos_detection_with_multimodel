"""Build a Word document containing blank templates for all 17 tables.

Source paper: A Two-Stage Hybrid Feature Selection and Ensemble ...
(Health Science Reports, 2026)

Headers, merged groups, row labels, units and model/feature names are kept.
All measured / result cells are intentionally left empty for later data entry.
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "docs" / "All_Tables.docx"

A4_W = Cm(21.0)
A4_H = Cm(29.7)
MARGIN = Cm(1.27)

HEADER_FILL = "D9E2F3"
GROUP_FILL = "BDD7EE"
LABEL_FILL = "EDEDED"

# Tables that must start on a fresh page.
PAGE_BREAK_BEFORE = {1, 3, 4, 5, 6, 8, 10, 11, 13, 15, 16}

# Tables that share a page: (first, second).
TABLE_PAIRS = [(1, 2), (6, 7), (8, 9), (11, 12), (13, 14), (16, 17)]
PAIR_FIRST = {first for first, _ in TABLE_PAIRS}
PAIR_SECOND = {second for _, second in TABLE_PAIRS}


# --------------------------------------------------------------------------
# Low-level helpers
# --------------------------------------------------------------------------
def setup_section(section, landscape=True):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = A4_H
    section.page_height = A4_W
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, size=9, align="center", fill=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0

    for index, line in enumerate(str(text).split("\n")):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Calibri"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), "Calibri"
        )

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        shade(cell, fill)


def repeat_header(table, row_indexes):
    for index in row_indexes:
        tr_pr = table.rows[index]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def make_table(doc, n_rows, n_cols, widths_in, size=9):
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False

    for column, width in zip(table.columns, widths_in):
        column.width = Cm(width * 2.54)
        for cell in column.cells:
            cell.width = Cm(width * 2.54)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            set_cell(cell, "", size=size)
    return table


def merge(t, r1, c1, r2, c2, text="", **kwargs):
    merged = t.cell(r1, c1).merge(t.cell(r2, c2))
    set_cell(merged, text, **kwargs)
    return merged


def add_caption(doc, number, caption):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(f"Table {number}. {caption}")
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = "Calibri"
    return paragraph


def begin_table_page(doc, number, caption, landscape=False):
    paragraph = add_caption(doc, number, caption)
    if number in PAGE_BREAK_BEFORE:
        paragraph.paragraph_format.page_break_before = True


def shrink_table(table, size=8, spacing=0):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(spacing)
                paragraph.paragraph_format.space_after = Pt(spacing)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(size)


def keep_rows_together(table, include_last=False):
    rows = table.rows if include_last else table.rows[:-1]
    for row in rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


# --------------------------------------------------------------------------
# Table builders
# --------------------------------------------------------------------------
def table_01(doc):
    caption = "Table of comparisons for existing works."
    begin_table_page(doc, 1, caption, landscape=True)
    headers = ["References", "Dataset", "Attributes used", "Method", "Limitations"]
    t = make_table(doc, 13, 5, [1.4, 1.8, 1.4, 1.4, 4.6])
    for col, text in enumerate(headers):
        set_cell(t.cell(0, col), text, bold=True, fill=HEADER_FILL)
    for row in range(1, 13):
        for col in range(5):
            set_cell(t.cell(row, col), "")
    repeat_header(t, [0])


def table_02(doc):
    caption = "Dataset distribution before and after oversampling."
    begin_table_page(doc, 2, caption)
    t = make_table(doc, 4, 6, [1.21] * 6)
    merge(t, 0, 0, 0, 5, "PCOS 541 patients' dataset", bold=True, fill=GROUP_FILL)
    merge(t, 1, 0, 1, 2, "Before oversampling", bold=True, fill=HEADER_FILL)
    merge(t, 1, 3, 1, 5, "After oversampling", bold=True, fill=HEADER_FILL)
    for col, text in enumerate(
        ["Training", "Testing", "Total", "Training", "Testing", "Total"]
    ):
        set_cell(t.cell(2, col), text, bold=True, fill=HEADER_FILL)
    for col in range(6):
        set_cell(t.cell(3, col), "")
    repeat_header(t, [0, 1, 2])


def table_03(doc):
    caption = (
        "Feature selection breakdown: From Chi-Square filtering to wrapper "
        "model refinement."
    )
    begin_table_page(doc, 3, caption, landscape=True)
    features = [
        "Age (years)",
        "Weight (kg)",
        "Height (cm)",
        "BMI",
        "Blood group",
        "RR (breaths/min)",
        "Hb (g/dL)",
        "Cycle (R/I)",
        "Cycle length (days)",
        "Marriage status (years)",
        "Pregnant (Y/N)",
        "No. of abortions",
        "II beta-HCG (mIU/mL)",
        "FSH (mIU/mL)",
        "Hip (inch)",
        "Waist (inch)",
        "AMH (ng/mL)",
        "Vit D3 (ng/mL)",
        "Weight gain (Y/N)",
        "Hair growth (Y/N)",
        "Skin darkening (Y/N)",
        "Hair loss (Y/N)",
        "Pimples (Y/N)",
        "Fast food (Y/N)",
        "Reg. exercise (Y/N)",
        "Follicle (L)",
        "Follicle (R)",
        "Avg. F size (L) (mm)",
        "Avg. F size (R) (mm)",
        "Endometrium (mm)",
    ]
    t = make_table(doc, 3 + len(features), 13, [2.0] + [0.72] * 12)
    merge(
        t,
        0,
        0,
        2,
        0,
        "Stage 1\nPreselected 30 features by \u03c72",
        bold=True,
        align="left",
        fill=GROUP_FILL,
    )
    merge(t, 0, 1, 0, 12, "Stage 2", bold=True, fill=GROUP_FILL)
    merge(
        t,
        1,
        1,
        1,
        4,
        "Selected in top 20 by wrapper model",
        bold=True,
        fill=HEADER_FILL,
    )
    merge(
        t,
        1,
        5,
        1,
        8,
        "Selected in top 16 by wrapper model",
        bold=True,
        fill=HEADER_FILL,
    )
    merge(
        t,
        1,
        9,
        1,
        12,
        "Selected in top 12 by wrapper model",
        bold=True,
        fill=HEADER_FILL,
    )
    for group in range(3):
        for offset, name in enumerate(["CB", "XGB", "LR", "RF"]):
            set_cell(
                t.cell(2, 1 + group * 4 + offset),
                name,
                bold=True,
                fill=HEADER_FILL,
            )
    for index, feature in enumerate(features, start=3):
        set_cell(t.cell(index, 0), feature, align="left")
        for col in range(1, 13):
            set_cell(t.cell(index, col), "")
    shrink_table(t, size=8, spacing=0)
    repeat_header(t, [0, 1, 2])


def table_04(doc):
    caption = "Final selected features using RFE."
    begin_table_page(doc, 4, caption)
    features = [
        "Follicle (R)",
        "Follicle (L)",
        "Hair growth (Y/N)",
        "Cycle (R/I)",
        "Weight gain (Y/N)",
        "Skin darkening (Y/N)",
        "Cycle length (days)",
        "Pimples (Y/N)",
        "Avg. F size (L) (mm)",
        "Marriage status (years)",
        "Avg. F size (R) (mm)",
        "RR (breaths/min)",
        "Fast food (Y/N)",
        "Weight (kg)",
        "FSH (mIU/mL)",
        "AMH (ng/mL)",
    ]
    t = make_table(doc, 1 + len(features), 2, [1.2, 6.0])
    for col, text in enumerate(["Sl. No.", "CB 16"]):
        set_cell(t.cell(0, col), text, bold=True, fill=HEADER_FILL)
    for index, feature in enumerate(features, start=1):
        set_cell(t.cell(index, 0), str(index))
        set_cell(t.cell(index, 1), feature, align="left")
    repeat_header(t, [0])


def table_05(doc):
    caption = "Statistical summary of common features."
    begin_table_page(doc, 5, caption, landscape=True)
    features = [
        "RR (breaths/min)",
        "Cycle (R/I)",
        "Cycle length (days)",
        "Marriage status (years)",
        "FSH (mIU/mL)",
        "AMH (ng/mL)",
        "Weight gain (Y/N)",
        "Hair growth (Y/N)",
        "Skin darkening (Y/N)",
        "Pimples (Y/N)",
        "Fast food (Y/N)",
        "Follicle (L)",
        "Follicle (R)",
        "Avg. F size (L) (mm)",
        "Avg. F size (R) (mm)",
        "Endometrium (mm)",
    ]
    headers = [
        "Feature",
        "\u03b2",
        "\u03b2-CI-lower",
        "\u03b2-CI-upper",
        "Y-correlation",
        "p",
        "Y1-mean",
        "Y0-mean",
        "All mean",
        "All Std",
    ]
    t = make_table(doc, 1 + len(features), 10, [1.9] + [0.97] * 9)
    for col, text in enumerate(headers):
        set_cell(t.cell(0, col), text, bold=True, fill=HEADER_FILL)
    for index, feature in enumerate(features, start=1):
        set_cell(t.cell(index, 0), feature, align="left")
        for col in range(1, 10):
            set_cell(t.cell(index, col), "")
    repeat_header(t, [0])


def table_06(doc):
    caption = "Software libraries and versions used in the analysis."
    begin_table_page(doc, 6, caption)
    libraries = [
        ("pandas", "LightGBM"),
        ("numpy", "XGBoost"),
        ("scikit-learn", "SHAP"),
        ("imbalanced-learn", "SciPy"),
        ("matplotlib", "statsmodels"),
        ("seaborn", "LIME"),
    ]
    t = make_table(doc, 1 + len(libraries), 4, [1.7, 1.1, 1.7, 1.1])
    for col, text in enumerate(["Library", "Version", "Library", "Version"]):
        set_cell(t.cell(0, col), text, bold=True, fill=HEADER_FILL)
    for row, (lib_a, lib_b) in enumerate(libraries, start=1):
        set_cell(t.cell(row, 0), lib_a, align="left")
        set_cell(t.cell(row, 1), "")
        set_cell(t.cell(row, 2), lib_b, align="left")
        set_cell(t.cell(row, 3), "")
    repeat_header(t, [0])


def table_07(doc):
    caption = "System environment configuration."
    begin_table_page(doc, 7, caption)
    t = make_table(doc, 5, 2, [2.1, 5.1])
    set_cell(t.cell(0, 0), "Resources", bold=True, fill=HEADER_FILL)
    set_cell(t.cell(0, 1), "Details", bold=True, fill=HEADER_FILL)
    for row, resource in enumerate(
        ["CPU", "GPU", "RAM", "Experimental tools"], start=1
    ):
        set_cell(t.cell(row, 0), resource, align="left")
        set_cell(t.cell(row, 1), "")
    repeat_header(t, [0])


def _wrapper_header(t, models):
    merge(t, 0, 0, 2, 0, "Model", bold=True, fill=GROUP_FILL)
    merge(
        t,
        0,
        1,
        0,
        12,
        "Wrapper model used in RFE",
        bold=True,
        fill=GROUP_FILL,
    )
    for group in range(3):
        merge(
            t,
            1,
            1 + group * 4,
            1,
            4 + group * 4,
            f"Using {[20, 16, 12][group]} features",
            bold=True,
            fill=HEADER_FILL,
        )
        for offset, name in enumerate(["LR (%)", "XGB (%)", "RF (%)", "CB (%)"]):
            set_cell(
                t.cell(2, 1 + group * 4 + offset),
                name,
                bold=True,
                fill=HEADER_FILL,
            )
    for row, model in enumerate(models, start=3):
        set_cell(t.cell(row, 0), model, align="left")
        for col in range(1, 13):
            set_cell(t.cell(row, col), "")
    repeat_header(t, [0, 1, 2])


def table_08(doc):
    caption = (
        "Model's accuracy of different ML models on various feature sets "
        "selected at the final stage of the two-stage feature selection."
    )
    begin_table_page(doc, 8, caption, landscape=True)
    models = ["LR", "DT", "RF", "XGB", "CB", "AB", "SVM", "KNN", "MLP", "LDA"]
    t = make_table(doc, 3 + len(models), 13, [1.15] + [0.79] * 12)
    _wrapper_header(t, models)


def table_09(doc):
    caption = (
        "Model's accuracy of stacking ensemble classifiers on different "
        "feature sets selected at the final stage of two-stage feature "
        "selection."
    )
    begin_table_page(doc, 9, caption, landscape=True)
    models = [
        "XGB + MLP",
        "XGB + CB",
        "LR + RF",
        "LR + AB + SVM",
        "SVM + MLP + LDA",
        "XGB + AB + MLP",
        "XGB + AB + SVM",
        "XGB + CB + KNN",
    ]
    t = make_table(doc, 3 + len(models), 13, [1.55] + [0.76] * 12)
    _wrapper_header(t, models)


def table_10(doc):
    caption = (
        "Statistical analysis data for various ML models with the most "
        "important four features."
    )
    begin_table_page(doc, 10, caption, landscape=True)
    feature_pairs = [
        ("Follicle (R)", "Cycle (R/I)"),
        ("Follicle (L)", "Hair growth (Y/N)"),
        ("Weight (kg)", "RR (breaths/min)"),
        ("Cycle length (days)", "Marriage status (years)"),
        ("FSH (mIU/mL)", "Avg. F size (L) (mm)"),
        ("Avg. F size (R) (mm)", "AMH (ng/mL)"),
        ("Weight gain (Y/N)", "Skin darkening (Y/N)"),
    ]
    classifiers = ["LR", "AB", "DT", "LGBM", "RF", "LDA", "XGB", "MLP"]
    metrics = [
        "\u03b2",
        "Y-correlation",
        "p",
        "Y1-mean",
        "Y0-mean",
        "All mean",
        "All Std",
    ]
    rows_per_band = 2 + len(metrics)
    n_rows = rows_per_band * len(feature_pairs)
    t = make_table(doc, n_rows, 17, [0.85] + [0.61] * 16, size=6)

    for band, (left, right) in enumerate(feature_pairs):
        base = band * rows_per_band
        set_cell(t.cell(base, 0), "Feature", bold=True, fill=GROUP_FILL, size=6)
        merge(
            t,
            base,
            1,
            base,
            8,
            left,
            bold=True,
            fill=GROUP_FILL,
            size=6,
        )
        merge(
            t,
            base,
            9,
            base,
            16,
            right,
            bold=True,
            fill=GROUP_FILL,
            size=6,
        )

        classifier_row = base + 1
        set_cell(
            t.cell(classifier_row, 0),
            "Classifier",
            bold=True,
            fill=HEADER_FILL,
            size=6,
        )
        for half in range(2):
            for offset, name in enumerate(classifiers):
                set_cell(
                    t.cell(classifier_row, 1 + half * 8 + offset),
                    name,
                    bold=True,
                    fill=HEADER_FILL,
                    size=6,
                )

        for offset, metric in enumerate(metrics, start=2):
            row = base + offset
            set_cell(t.cell(row, 0), metric, bold=True, size=6, fill=LABEL_FILL)
            for col in range(1, 17):
                set_cell(t.cell(row, col), "", size=6)

    for band in range(len(feature_pairs)):
        base = band * rows_per_band
        repeat_header(t, [base, base + 1])


def table_11(doc):
    caption = (
        "Key findings ML model versus final feature set (16 features contains)."
    )
    begin_table_page(doc, 11, caption)
    t = make_table(doc, 8, 5, [2.0] + [1.3] * 4)
    merge(t, 0, 0, 1, 0, "Matrices", bold=True, fill=GROUP_FILL)
    merge(
        t,
        0,
        1,
        0,
        4,
        "16 Features and different wrapper models used in RFE",
        bold=True,
        fill=GROUP_FILL,
    )
    for col, text in enumerate(["LR (%)", "XGB (%)", "RF (%)", "CB (%)"], start=1):
        set_cell(t.cell(1, col), text, bold=True, fill=HEADER_FILL)
    for row, label in enumerate(
        [
            "1st Rank accuracy",
            "2nd Rank accuracy",
            "3rd Rank accuracy",
            "4th Rank accuracy",
            "Minimum accuracy",
            "Average accuracy",
        ],
        start=2,
    ):
        set_cell(t.cell(row, 0), label, align="left")
        for col in range(1, 5):
            set_cell(t.cell(row, col), "")
    repeat_header(t, [0, 1])


def _performance_table(doc, number, caption, models):
    begin_table_page(doc, number, caption)
    headers = [
        "Model",
        "Accuracy (%)",
        "Precision (%)",
        "Recall (%)",
        "F1 score (%)",
    ]
    t = make_table(doc, 1 + len(models), 5, [2.2] + [1.25] * 4)
    for col, text in enumerate(headers):
        set_cell(t.cell(0, col), text, bold=True, fill=HEADER_FILL)
    for row, model in enumerate(models, start=1):
        set_cell(t.cell(row, 0), model, align="left")
        for col in range(1, 5):
            set_cell(t.cell(row, col), "")
    repeat_header(t, [0])


def table_12(doc):
    caption = "Model's performance of different ML models using the proposed feature set."
    models = ["LR", "DT", "RF", "XGB", "CB", "AB", "SVM", "KNN", "MLP", "LDA"]
    _performance_table(doc, 12, caption, models)


def table_13(doc):
    caption = (
        "Model's performance of stacking ensemble classifiers using the "
        "proposed feature set."
    )
    models = [
        "XGB + MLP",
        "XGB + CB",
        "LR + RF",
        "DT + RF",
        "LR + AB + SVM",
        "SVM + MLP + LDA",
        "XGB + AB + MLP",
        "XGB + AB + SVM",
        "XGB + CB + KNN",
    ]
    _performance_table(doc, 13, caption, models)


def table_14(doc):
    caption = (
        "Overfitting analysis across classifiers with and without feature "
        "selection."
    )
    begin_table_page(doc, 14, caption, landscape=True)
    models = [
        "LR",
        "DT",
        "RF",
        "AB",
        "XGB",
        "CB",
        "SVM",
        "KNN",
        "GNB",
        "MLP",
        "LDA",
        "XGB + MLP",
    ]
    t = make_table(doc, 2 + len(models), 7, [1.15] + [1.55] * 6)
    merge(t, 0, 0, 1, 0, "Model", bold=True, fill=GROUP_FILL)
    merge(
        t,
        0,
        1,
        0,
        3,
        "Overfitting analysis with selected features",
        bold=True,
        fill=GROUP_FILL,
        size=8,
    )
    merge(
        t,
        0,
        4,
        0,
        6,
        "Overfitting analysis without feature selection",
        bold=True,
        fill=GROUP_FILL,
        size=8,
    )
    sub = [
        "Train accuracy (%)",
        "Test accuracy (%)",
        "Overfitting (%)",
        "Train accuracy (%)",
        "Test accuracy (%)",
        "Overfitting (%)",
    ]
    for col, text in enumerate(sub, start=1):
        set_cell(t.cell(1, col), text, bold=True, fill=HEADER_FILL, size=8)
    for row, model in enumerate(models, start=2):
        set_cell(t.cell(row, 0), model, align="left")
        for col in range(1, 7):
            set_cell(t.cell(row, col), "")
    repeat_header(t, [0, 1])


def table_15(doc):
    caption = (
        "Performance metrics of individual and ensemble machine learning "
        "models using 10-fold stratified cross-validation."
    )
    models = [
        "CB",
        "XGB + CB + KNN (Soft Voting)",
        "SVM + MLP + LDA (Soft Voting)",
        "LR + AdaBoost + SVM (Soft Voting)",
        "XGB",
        "XGB + CB (Soft Voting)",
        "XGB + MLP (Soft Voting)",
        "XGB + AdaBoost + MLP (Soft Voting)",
        "XGB + AdaBoost + SVM (Soft Voting)",
        "MLP",
        "LR",
        "SVM",
        "AdaBoost",
        "LDA",
        "KNN",
        "DT",
    ]
    _performance_table(doc, 15, caption, models)


def table_16(doc):
    caption = "Computational complexity on different stage of proposed method."
    begin_table_page(doc, 16, caption)
    t = make_table(doc, 4, 3, [2.3, 2.6, 2.4])
    for col, text in enumerate(
        ["Stage", "Operation", "Theoretical complexity"]
    ):
        set_cell(t.cell(0, col), text, bold=True, fill=HEADER_FILL)
    rows = [
        ("Chi-Square filtering", "Univariate dependence testing"),
        ("RFE with CatBoost", "Iterative elimination and retraining"),
        ("Model training (Final Ensemble)", "Training and evaluation"),
    ]
    for row, (stage, operation) in enumerate(rows, start=1):
        set_cell(t.cell(row, 0), stage, align="left")
        set_cell(t.cell(row, 1), operation, align="left")
        set_cell(t.cell(row, 2), "")
    repeat_header(t, [0])


def table_17(doc):
    caption = "A comparison analysis between the proposed method and recent works."
    begin_table_page(doc, 17, caption)
    headers = [
        "Title",
        "Method",
        "XAI",
        "Accuracy (%)",
        "Precision (%)",
        "Recall (%)",
        "F1 score (%)",
    ]
    t = make_table(doc, 10, 7, [2.0, 1.2, 1.1, 0.85, 0.85, 0.75, 0.75])
    for col, text in enumerate(headers):
        set_cell(t.cell(0, col), text, bold=True, fill=HEADER_FILL, size=8)
    for row in range(1, 9):
        for col in range(7):
            set_cell(t.cell(row, col), "", size=8)
    set_cell(t.cell(9, 0), "Proposed method", bold=True, align="left", size=8)
    for col in range(1, 7):
        set_cell(t.cell(9, col), "", size=8)
    repeat_header(t, [0])


# --------------------------------------------------------------------------
# Document assembly
# --------------------------------------------------------------------------
def add_title_page(doc):
    section = doc.sections[0]
    setup_section(section, landscape=False)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)
    run = title.add_run("All Tables \u2014 Blank Templates")
    run.bold = True
    run.font.size = Pt(24)
    run.font.name = "Calibri"

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "A Two-Stage Hybrid Feature Selection and Ensemble Learning "
        "Framework for PCOS Prediction"
    )
    run.font.size = Pt(12)
    run.italic = True

    source = doc.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = source.add_run("Health Science Reports, 2026")
    run.font.size = Pt(10)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "17 tables \u2022 Headers, merged cells, units and row labels are "
        "preserved.\nAll measured/result cells are intentionally blank for "
        "data entry."
    )
    run.font.size = Pt(10)
    run.font.color.rgb = None

    doc.add_paragraph()
    list_heading = doc.add_paragraph()
    run = list_heading.add_run("Contents")
    run.bold = True
    run.font.size = Pt(12)

    captions = [
        "Table of comparisons for existing works",
        "Dataset distribution before and after oversampling",
        "Feature selection breakdown",
        "Final selected features using RFE",
        "Statistical summary of common features",
        "Software libraries and versions",
        "System environment configuration",
        "Model accuracy on different feature sets",
        "Stacking ensemble accuracy on different feature sets",
        "Statistical analysis data for various ML models",
        "Key findings: ML model versus final feature set",
        "Model performance using the proposed feature set",
        "Stacking ensemble performance",
        "Overfitting analysis",
        "10-fold stratified cross-validation metrics",
        "Computational complexity by stage",
        "Comparison with recent works",
    ]
    for index, text in enumerate(captions, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Cm(0.6)
        run = paragraph.add_run(f"Table {index}. {text}")
        run.font.size = Pt(9)


def main():
    global OUT_PATH
    if len(sys.argv) > 1:
        OUT_PATH = Path(sys.argv[1])

    doc = Document()
    add_title_page(doc)

    builders = [
        table_01,
        table_02,
        table_03,
        table_04,
        table_05,
        table_06,
        table_07,
        table_08,
        table_09,
        table_10,
        table_11,
        table_12,
        table_13,
        table_14,
        table_15,
        table_16,
        table_17,
    ]
    for builder in builders:
        builder(doc)

    for index, table in enumerate(doc.tables, start=1):
        if index in PAIR_FIRST:
            keep_rows_together(table, include_last=True)
        elif index in PAIR_SECOND:
            keep_rows_together(table, include_last=False)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")
    print(f"Tables created: {len(doc.tables)}")
    for index, table in enumerate(doc.tables, start=1):
        print(f"  Table {index:>2}: {len(table.rows)} rows x {len(table.columns)} cols")


if __name__ == "__main__":
    main()
